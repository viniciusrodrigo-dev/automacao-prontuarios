from pathlib import Path
from uuid import uuid4
import os
import re
import json
import subprocess
import sys
import shutil
from datetime import date, datetime

from flask import Flask, render_template, request, jsonify
from openpyxl import load_workbook

pasta_projeto = Path(__file__).resolve().parent
app = Flask(__name__, template_folder="web/templates", static_folder="web/static")
app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024

pasta_uploads = pasta_projeto / "entrada_web"
pasta_uploads.mkdir(parents=True, exist_ok=True)
caminho_planilha = pasta_projeto / "revisao_funcionarios.xlsx"
COLUNAS = {
    "arquivo": "Arquivo PDF",
    "pagina": "Página",
    "nome": "Nome",
    "data_admissao": "Data de admissão",
    "cargo": "Cargo",
    "projeto": "Projeto",
    "status": "Status",
    "pendencias": "Pendências",
    "revisado": "Revisado manualmente",
}
CAMPOS_OBRIGATORIOS = {
    "nome": "Nome",
    "data_admissao": "Data de admissão",
    "cargo": "Cargo",
    "projeto": "Projeto",
}


def validar_lote(id_lote):
    if not isinstance(id_lote, str) or not re.fullmatch(r"[a-f0-9]{32}", id_lote):
        return None
    pasta_lote = pasta_uploads / id_lote
    if not pasta_lote.is_dir() or pasta_lote.parent != pasta_uploads:
        return None
    return pasta_lote


def ler_manifesto(pasta_lote):
    caminho = pasta_lote / "lote.json"
    if not caminho.exists():
        # Compatibilidade com lotes processados antes da criação do manifesto próprio.
        caminho = pasta_projeto / "lote_atual.json"
    try:
        dados = json.loads(caminho.read_text(encoding="utf-8"))
        return dados if isinstance(dados, list) else []
    except (OSError, json.JSONDecodeError):
        return []


def abrir_planilha():
    if not caminho_planilha.exists():
        return None, None
    workbook = load_workbook(caminho_planilha)
    if "Funcionários" not in workbook.sheetnames:
        workbook.close()
        return None, None
    planilha = workbook["Funcionários"]
    colunas = {
        str(c.value).strip(): i
        for i, c in enumerate(planilha[1], start=1)
        if c.value is not None
    }
    if any(nome not in colunas for nome in COLUNAS.values()):
        workbook.close()
        return None, None
    return workbook, (planilha, colunas)


def texto_celula(valor):
    if valor is None:
        return ""
    if isinstance(valor, (datetime, date)):
        return valor.strftime("%d/%m/%Y")
    return str(valor).strip()


@app.route("/")
def inicio():
    return render_template("index.html")


@app.route("/api/upload", methods=["POST"])
def receber_pdfs():
    arquivos = [a for a in request.files.getlist("arquivos") if a and a.filename]
    if not arquivos:
        return jsonify({"sucesso": False, "mensagem": "Selecione pelo menos um PDF."}), 400
    for arquivo in arquivos:
        if Path(arquivo.filename).suffix.lower() != ".pdf":
            return jsonify({"sucesso": False, "mensagem": f"'{Path(arquivo.filename).name}' não é PDF."}), 400

    id_lote = uuid4().hex
    pasta_lote = pasta_uploads / id_lote
    pasta_lote.mkdir(parents=True, exist_ok=False)
    salvos = []
    try:
        for arquivo in arquivos:
            original = Path(arquivo.filename).name
            armazenado = f"{uuid4().hex}.pdf"
            arquivo.save(pasta_lote / armazenado)
            salvos.append({"nome_original": original, "nome_armazenado": armazenado})
    except Exception:
        shutil.rmtree(pasta_lote, ignore_errors=True)
        app.logger.exception("Erro ao salvar PDFs")
        return jsonify({"sucesso": False, "mensagem": "Não foi possível salvar todos os PDFs."}), 500
    return jsonify({"sucesso": True, "mensagem": "PDFs recebidos.", "lote": id_lote,
                    "quantidade": len(salvos), "arquivos": salvos})


@app.route("/api/processar", methods=["POST"])
def processar_pdfs():
    dados = request.get_json(silent=True) or {}
    id_lote = dados.get("lote", "")
    pasta_lote = validar_lote(id_lote)
    if pasta_lote is None:
        return jsonify({"sucesso": False, "mensagem": "Lote inválido ou não encontrado."}), 400
    if not list(pasta_lote.glob("*.pdf")):
        return jsonify({"sucesso": False, "mensagem": "Não há PDFs neste lote."}), 400

    ambiente = os.environ.copy()
    ambiente["pasta_entrada"] = str(pasta_lote.resolve())
    ambiente["caminho_lote"] = str((pasta_lote / "lote.json").resolve())
    try:
        resultado = subprocess.run(
            [sys.executable, "-m", "src.processar_documentos"],
            cwd=str(pasta_projeto), env=ambiente, capture_output=True,
            text=True, encoding="utf-8", errors="replace", timeout=600, check=False
        )
    except subprocess.TimeoutExpired:
        return jsonify({"sucesso": False, "mensagem": "O OCR excedeu 10 minutos."}), 504
    except Exception:
        app.logger.exception("Falha ao iniciar OCR")
        return jsonify({"sucesso": False, "mensagem": "Não foi possível iniciar o OCR."}), 500

    saida = resultado.stdout or ""
    erros = resultado.stderr or ""
    if resultado.returncode != 0:
        return jsonify({"sucesso": False, "mensagem": "O OCR encontrou um erro.",
                        "detalhes": (erros or saida)[-5000:]}), 500
    return jsonify({"sucesso": True, "mensagem": "Processamento OCR concluído.",
                    "detalhes": saida[-5000:], "lote": id_lote})


@app.route("/api/funcionarios", methods=["GET"])
def consultar_funcionarios():
    id_lote = request.args.get("lote", "")
    pasta_lote = validar_lote(id_lote)
    if pasta_lote is None:
        return jsonify({"sucesso": False, "mensagem": "Lote inválido ou não encontrado."}), 400
    manifesto = ler_manifesto(pasta_lote)
    chaves = {(str(i.get("arquivo_pdf", "")), str(i.get("pagina", ""))) for i in manifesto}
    if not chaves:
        return jsonify({"sucesso": True, "funcionarios": [], "mensagem": "Nenhum registro neste lote."})
    workbook, contexto = abrir_planilha()
    if workbook is None:
        return jsonify({"sucesso": False, "mensagem": "Planilha de revisão ausente ou com colunas inválidas."}), 500
    planilha, colunas = contexto
    funcionarios = []
    for linha in range(2, planilha.max_row + 1):
        arquivo = texto_celula(planilha.cell(linha, colunas[COLUNAS["arquivo"]]).value)
        pagina = texto_celula(planilha.cell(linha, colunas[COLUNAS["pagina"]]).value)
        if (arquivo, pagina) not in chaves:
            continue
        funcionarios.append({
            "arquivo_pdf": arquivo, "pagina": pagina,
            "nome": texto_celula(planilha.cell(linha, colunas[COLUNAS["nome"]]).value),
            "data_admissao": texto_celula(planilha.cell(linha, colunas[COLUNAS["data_admissao"]]).value),
            "cargo": texto_celula(planilha.cell(linha, colunas[COLUNAS["cargo"]]).value),
            "projeto": texto_celula(planilha.cell(linha, colunas[COLUNAS["projeto"]]).value),
            "status": texto_celula(planilha.cell(linha, colunas[COLUNAS["status"]]).value),
        })
    workbook.close()
    return jsonify({"sucesso": True, "funcionarios": funcionarios})


@app.route("/api/funcionarios/salvar", methods=["POST"])
def salvar_funcionarios():
    dados = request.get_json(silent=True) or {}
    id_lote = dados.get("lote", "")
    registros = dados.get("funcionarios", [])
    pasta_lote = validar_lote(id_lote)
    if pasta_lote is None or not isinstance(registros, list):
        return jsonify({"sucesso": False, "mensagem": "Lote ou dados inválidos."}), 400
    manifesto = ler_manifesto(pasta_lote)
    chaves_lote = {(str(i.get("arquivo_pdf", "")), str(i.get("pagina", ""))) for i in manifesto}
    workbook, contexto = abrir_planilha()
    if workbook is None:
        return jsonify({"sucesso": False, "mensagem": "Planilha de revisão ausente ou inválida."}), 500
    planilha, colunas = contexto
    indices = {}
    for linha in range(2, planilha.max_row + 1):
        chave = (texto_celula(planilha.cell(linha, colunas[COLUNAS["arquivo"]]).value),
                 texto_celula(planilha.cell(linha, colunas[COLUNAS["pagina"]]).value))
        if chave in chaves_lote:
            indices[chave] = linha
    salvos = 0
    for registro in registros:
        if not isinstance(registro, dict):
            continue
        chave = (str(registro.get("arquivo_pdf", "")), str(registro.get("pagina", "")))
        linha = indices.get(chave)
        if chave not in chaves_lote or linha is None:
            continue
        for campo, coluna_nome in CAMPOS_OBRIGATORIOS.items():
            valor = str(registro.get(campo, "") or "").strip()
            planilha.cell(linha, colunas[COLUNAS[campo]]).value = valor or "PREENCHER"
        faltantes = [nome for campo, nome in CAMPOS_OBRIGATORIOS.items()
                     if not str(planilha.cell(linha, colunas[COLUNAS[campo]]).value or "").strip()
                     or str(planilha.cell(linha, colunas[COLUNAS[campo]]).value).strip().upper() == "PREENCHER"]
        planilha.cell(linha, colunas[COLUNAS["status"]]).value = "REVISAR" if faltantes else "PRONTO"
        planilha.cell(linha, colunas[COLUNAS["pendencias"]]).value = ", ".join(faltantes)
        planilha.cell(linha, colunas[COLUNAS["revisado"]]).value = "SIM"
        salvos += 1
    workbook.save(caminho_planilha)
    workbook.close()
    return jsonify({"sucesso": True, "mensagem": f"{salvos} registro(s) salvo(s).", "salvos": salvos})

@app.route("/api/gerar", methods=["POST"])
def gerar_documentos():
    dados = request.get_json(silent=True) or {}
    id_lote = dados.get("lote", "")

    pasta_lote = validar_lote(id_lote)

    if pasta_lote is None:
        return jsonify({
            "sucesso": False,
            "mensagem": "Lote inválido ou não encontrado."
        }), 400

    manifesto = ler_manifesto(pasta_lote)

    if not manifesto:
        return jsonify({
            "sucesso": False,
            "mensagem": "O lote não possui registros processados."
        }), 400

    ambiente = os.environ.copy()
    ambiente["caminho_lote"] = str(
        (pasta_lote / "lote.json").resolve()
    )

    try:
        resultado = subprocess.run(
            [sys.executable, "-m", "src.gerar_capas"],
            cwd=str(pasta_projeto),
            env=ambiente,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=600,
            check=False
        )

    except subprocess.TimeoutExpired:
        app.logger.exception(
            "Tempo limite excedido durante a geração do lote %s",
            id_lote
        )

        return jsonify({
            "sucesso": False,
            "mensagem": "A geração excedeu 10 minutos."
        }), 504

    except Exception as erro:
        app.logger.exception(
            "Falha ao iniciar a geração do lote %s",
            id_lote
        )

        return jsonify({
            "sucesso": False,
            "mensagem": "Não foi possível iniciar a geração.",
            "detalhes": str(erro)
        }), 500

    saida = resultado.stdout or ""
    erros = resultado.stderr or ""

    # Registra os detalhes no terminal onde o Flask está executando.
    app.logger.info(
        "Saída do gerador para o lote %s:\n%s",
        id_lote,
        saida[-5000:]
    )

    if erros:
        app.logger.error(
            "Erros do gerador para o lote %s:\n%s",
            id_lote,
            erros[-5000:]
        )

    detalhes = "\n".join(
        parte for parte in [
            saida.strip(),
            erros.strip()
        ] if parte
    )

    if resultado.returncode != 0:
        return jsonify({
            "sucesso": False,
            "mensagem": "Erro ao gerar documentos.",
            "detalhes": detalhes[-5000:],
            "codigo_saida": resultado.returncode
        }), 500

    if (
        "GERAÇÃO BLOQUEADA" in saida
        or "Nenhum funcionário válido" in saida
    ):
        return jsonify({
            "sucesso": False,
            "mensagem": "Geração bloqueada. Confira os campos obrigatórios.",
            "detalhes": detalhes[-5000:]
        }), 400

    return jsonify({
        "sucesso": True,
        "mensagem": "Geração concluída.",
        "detalhes": detalhes[-5000:]
    })

@app.errorhandler(413)
def arquivo_muito_grande(erro):
    return jsonify({"sucesso": False, "mensagem": "O envio ultrapassou o limite de 50 MB."}), 413


if __name__ == "__main__":
    app.run(debug=True)
