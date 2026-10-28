from pathlib import Path
from datetime import datetime, date
import json
import os
from openpyxl import load_workbook

from .preencher_capa import preencher_capa
from .preencher_declaracao import preencher_declaracao
from .relatorio import RelatorioProcessamento

pasta_projeto = Path(__file__).resolve().parent.parent
caminho_planilha = pasta_projeto / "revisao_funcionarios.xlsx"
caminho_lote = Path(os.environ.get("caminho_lote", str(pasta_projeto / "lote_atual.json")))
relatorio = RelatorioProcessamento(pasta_projeto)

if not caminho_planilha.exists():
    print("ERRO: A planilha de revisão não foi encontrada.")
    raise SystemExit(1)
if not caminho_lote.exists():
    print("ERRO: Não existe lote atual. Execute processar_documentos.py primeiro.")
    raise SystemExit(1)

try:
    lote = json.loads(caminho_lote.read_text(encoding="utf-8"))
except (OSError, json.JSONDecodeError) as erro:
    print(f"ERRO ao ler o lote atual: {erro}")
    raise SystemExit(1)

chaves_lote = {(str(i["arquivo_pdf"]), str(i["pagina"])) for i in lote}
if not chaves_lote:
    print("O lote atual está vazio.")
    raise SystemExit(0)

workbook = load_workbook(caminho_planilha)
if "Funcionários" not in workbook.sheetnames:
    workbook.close()
    print("ERRO: A aba 'Funcionários' não existe.")
    raise SystemExit(1)
planilha = workbook["Funcionários"]

colunas = {
    str(c.value).strip(): i
    for i, c in enumerate(planilha[1], start=1) if c.value is not None
}
necessarias = ["Arquivo PDF", "Página", "Nome", "Data de admissão", "Cargo", "Projeto", "Status", "Pendências"]
faltantes = [c for c in necessarias if c not in colunas]
if faltantes:
    workbook.close()
    print("ERRO: faltam colunas na planilha: " + ", ".join(faltantes))
    raise SystemExit(1)

funcionarios = []
pendencias_gerais = []
for linha in range(2, planilha.max_row + 1):
    arquivo = str(planilha.cell(linha, colunas["Arquivo PDF"]).value or "")
    pagina = str(planilha.cell(linha, colunas["Página"]).value or "")
    if (arquivo, pagina) not in chaves_lote:
        continue

    dados = {
        "nome": planilha.cell(linha, colunas["Nome"]).value,
        "data_admissao": planilha.cell(linha, colunas["Data de admissão"]).value,
        "cargo": planilha.cell(linha, colunas["Cargo"]).value,
        "projeto": planilha.cell(linha, colunas["Projeto"]).value,
        "data_demissao": ""
    }
    vazios = [
        campo for campo in ("nome", "data_admissao", "cargo", "projeto")
        if dados[campo] is None
        or not str(dados[campo]).strip()
        or str(dados[campo]).strip().upper() == "PREENCHER"
    ]
    if vazios:
        pendencias_gerais.append(f"Linha {linha}: {', '.join(vazios)}")
    else:
        funcionarios.append((linha, dados, arquivo, pagina))

if pendencias_gerais:
    workbook.close()
    print("\nGERAÇÃO BLOQUEADA. Corrija os campos do lote atual no Excel:")
    for p in pendencias_gerais:
        print("-", p)
    raise SystemExit(0)
if not funcionarios:
    workbook.close()
    print("Nenhum funcionário válido encontrado no lote atual.")
    raise SystemExit(0)

total_capa = total_capa_nao = total_decl = total_decl_nao = 0
print(f"\nGerando documentos para {len(funcionarios)} registro(s) do lote atual.\n")

for linha, dados, arquivo, pagina in funcionarios:
    data = dados["data_admissao"]
    if isinstance(data, (datetime, date)):
        dados["data_admissao"] = data.strftime("%d/%m/%Y")
    else:
        dados["data_admissao"] = str(data).strip()

    capa_ok = preencher_capa(dados)
    declaracao_ok = preencher_declaracao(dados)

    total_capa += int(bool(capa_ok))
    total_capa_nao += int(not capa_ok)
    total_decl += int(bool(declaracao_ok))
    total_decl_nao += int(not declaracao_ok)

    if capa_ok:
        planilha.cell(linha, colunas["Status"]).value = "CAPA CRIADA"
        planilha.cell(linha, colunas["Pendências"]).value = ""
    else:
        planilha.cell(linha, colunas["Status"]).value = "CAPA NÃO CRIADA"

    relatorio.registrar(
        arquivo_pdf=arquivo, pagina=pagina, dados=dados,
        resultado="CAPA CRIADA" if capa_ok else "CAPA NÃO CRIADA",
        detalhes="Processamento do lote atual; declaração "
                 + ("criada." if declaracao_ok else "não criada ou já existente.")
    )

workbook.save(caminho_planilha)
workbook.close()
print("\nPROCESSO CONCLUÍDO!")
print(f"Capas criadas: {total_capa}")
print(f"Capas não criadas (erro ou já existentes): {total_capa_nao}")
print(f"Declarações criadas: {total_decl}")
print(f"Declarações não criadas (erro ou já existentes): {total_decl_nao}")
print("Os registros antigos da planilha não pertencentes ao lote atual foram ignorados.")
