from pathlib import Path
from datetime import datetime, date
from openpyxl import load_workbook
from .preencher_capa import preencher_capa
from .preencher_declaracao import preencher_declaracao
import sys
import subprocess

PASTA = Path(__file__).resolve().parent.parent
PLANILHA = PASTA / "revisao_funcionarios.xlsx"

def data_texto(valor):
    if isinstance(valor, (datetime, date)):
        return valor.strftime("%d/%m/%Y")
    return str(valor or "").strip()

def carregar_funcionarios():
    if not PLANILHA.exists():
        print("Planilha de revisão não encontrada.")
        return []
    wb = load_workbook(PLANILHA, data_only=True)
    try:
        ws = wb["Funcionários"]
        colunas = {str(c.value).strip(): i for i, c in enumerate(ws[1], 1) if c.value}
        obrigatorias = ["Nome", "Data de admissão", "Cargo", "Projeto"]
        if any(c not in colunas for c in obrigatorias):
            print("A planilha não contém todas as colunas necessárias.")
            return []
        registros = []
        for linha in range(2, ws.max_row + 1):
            dados = {
                "nome": ws.cell(linha, colunas["Nome"]).value,
                "data_admissao": data_texto(ws.cell(linha, colunas["Data de admissão"]).value),
                "cargo": ws.cell(linha, colunas["Cargo"]).value,
                "projeto": ws.cell(linha, colunas["Projeto"]).value,
                "data_demissao": ""
            }
            if str(dados["nome"] or "").strip() and str(dados["nome"]).strip().upper() != "PREENCHER":
                registros.append((linha, dados))
        return registros
    finally:
        wb.close()

def main():
    while True:
        print("\n=== GERENCIADOR DE DOCUMENTOS ===")
        print("1 - Gerar documentos do lote atual")
        print("2 - Reimprimir documentos de um funcionário do histórico")
        print("3 - Abrir a planilha de revisão")
        print("0 - Sair")
        opcao = input("Escolha: ").strip()

        if opcao == "1":
            import subprocess, sys
            subprocess.run(
                [sys.executable, "-m", "src.gerar_capas"],
                cwd=PASTA,
                check=True
            )
        elif opcao == "2":
            registros = carregar_funcionarios()
            if not registros:
                print("Nenhum registro com nome encontrado.")
                continue
            termo = input("Digite parte do nome do funcionário: ").strip().casefold()
            encontrados = [(l, d) for l, d in registros if termo in str(d["nome"]).casefold()]
            if not encontrados:
                print("Nenhum funcionário encontrado.")
                continue
            for indice, (linha, dados) in enumerate(encontrados, 1):
                print(f"{indice} - {dados['nome']} | admissão: {dados['data_admissao']} | linha {linha}")
            try:
                escolha = int(input("Selecione o funcionário: "))
                if not 1 <= escolha <= len(encontrados):
                    raise ValueError
            except ValueError:
                print("Seleção inválida.")
                continue
            _, dados = encontrados[escolha - 1]
            print("1 - Capa do prontuário")
            print("2 - Declaração de alimentação")
            print("3 - Ambos")
            tipo = input("Escolha: ").strip()
            if tipo not in {"1", "2", "3"}:
                print("Opção inválida.")
                continue
            print("\nObservação: os geradores atuais não sobrescrevem arquivos existentes.")
            print("Se o arquivo já existir, mova-o para outra pasta antes de reimprimir.")
            if tipo in {"1", "3"}:
                ok = preencher_capa(dados)
                print("Capa criada." if ok else "Capa não criada (pode já existir ou ter ocorrido erro).")
            if tipo in {"2", "3"}:
                ok = preencher_declaracao(dados)
                print("Declaração criada." if ok else "Declaração não criada (pode já existir ou ter ocorrido erro).")
        elif opcao == "3":
            import os
            os.startfile(PLANILHA)
        elif opcao == "0":
            break
        else:
            print("Opção inválida.")

if __name__ == "__main__":
    main()
