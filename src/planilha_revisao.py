from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from pathlib import Path


class PlanilhaRevisao:

    def __init__(self, pasta_projeto):
        self.pasta_projeto = Path(pasta_projeto)

        self.caminho_arquivo = (
            self.pasta_projeto / "revisao_funcionarios.xlsx"
        )

        self.colunas = [
            "Arquivo PDF",
            "Página",
            "Nome",
            "Data de admissão",
            "Cargo",
            "Projeto",
            "Status",
            "Pendências",
            "Revisado manualmente"
        ]

    def criar_planilha(self):
        # Não recria uma planilha que já existe.
        if self.caminho_arquivo.exists():
            return

        workbook = Workbook()
        planilha = workbook.active
        planilha.title = "Funcionários"

        planilha.append(self.colunas)

        for celula in planilha[1]:
            celula.font = Font(
                bold=True,
                color="FFFFFF"
            )
            celula.fill = PatternFill(
                fill_type="solid",
                fgColor="1F4E78"
            )
            celula.alignment = Alignment(
                horizontal="center"
            )

        larguras = [30, 10, 35, 20, 35, 30, 20, 35, 25]

        for indice, largura in enumerate(larguras, start=1):
            letra = get_column_letter(indice)
            planilha.column_dimensions[letra].width = largura

        planilha.freeze_panes = "A2"

        planilha.auto_filter.ref = planilha.dimensions

        workbook.save(self.caminho_arquivo)

        print(
            f"Planilha criada: {self.caminho_arquivo}"
        )

    def registrar_funcionario(
        self,
        arquivo_pdf,
        pagina,
        dados,
        pendencias
    ):
        self.criar_planilha()

        workbook = load_workbook(self.caminho_arquivo)
        planilha = workbook["Funcionários"]

        # Procura um registro do mesmo PDF e da mesma página.
        linha_existente = None

        for linha in range(2, planilha.max_row + 1):
            mesmo_arquivo = (
                planilha.cell(linha, 1).value == arquivo_pdf
            )

            mesma_pagina = (
                str(planilha.cell(linha, 2).value) == str(pagina)
            )

            if mesmo_arquivo and mesma_pagina:
                linha_existente = linha
                break

        # Se já existe, não sobrescreve os dados.
        # Isso protege as correções feitas manualmente.
        if linha_existente is not None:
            workbook.close()
            print(
                f"Registro preservado: {arquivo_pdf}, "
                f"página {pagina}."
            )
            return

        campos_obrigatorios = [
            "nome",
            "data_admissao",
            "cargo",
            "projeto"
        ]

        pendencias = list(pendencias)

        status = "REVISAR" if pendencias else "PRONTO"

        valores = [
            dados.get("nome") or "PREENCHER",
            dados.get("data_admissao") or "PREENCHER",
            dados.get("cargo") or "PREENCHER",
            dados.get("projeto") or "PREENCHER"
        ]

        # Confere também os valores que realmente serão registrados.
        for campo, valor in zip(campos_obrigatorios, valores):
            if valor == "PREENCHER" and campo not in pendencias:
                pendencias.append(campo)

        status = "REVISAR" if pendencias else "PRONTO"

        planilha.append([
            arquivo_pdf,
            pagina,
            *valores,
            status,
            ", ".join(pendencias),
            "NÃO"
        ])

        linha = planilha.max_row

        for coluna in range(1, len(self.colunas) + 1):
            celula = planilha.cell(linha, coluna)

            if status == "REVISAR":
                celula.fill = PatternFill(
                    fill_type="solid",
                    fgColor="FFF2CC"
                )

        planilha.auto_filter.ref = planilha.dimensions

        workbook.save(self.caminho_arquivo)
        workbook.close()

        print(
            f"Funcionário registrado: {arquivo_pdf}, "
            f"página {pagina}."
        )